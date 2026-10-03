"""
core/woff.py
Write-Off report automation.

Handles loading, updating, and generating the SP Madrid
CBS Auto Write-Off Status Report (WOFF STATUS + FIELD sheets).
Same logic as the Recovery report but adapted for the Write-Off
column layout.
"""
from __future__ import annotations

import io
from typing import Optional

import openpyxl
import pandas as pd

from core.file_io import decrypt_workbook, encrypt_workbook
from core.remarks import (
    update_status_remarks,
    update_fv_remarks,
    select_daily_action_code,
    merge_repo_ai_into_fv,
    FV_PLACEHOLDER,
)
from core.status_lookup import get_lookup

# ---------------------------------------------------------------------------
# Column alias lists for fuzzy matching
# ---------------------------------------------------------------------------
WOFF_PN_ALIASES       = ["PN#", "PN NO", "PN", "ACCOUNT NO", "ACCOUNT NUMBER"]
WOFF_NAME_ALIASES     = ["ACCOUNT NAME", "CUST_NAME", "NAME", "CLIENT NAME"]
WOFF_AC_ALIASES       = ["ACTION CODE", "DAILY STATUS", "CURRENT STATUS"]
WOFF_ACCT_STATUS_ALIASES = ["ACCOUNT STATUS", "COLLECTION ACTIVITIES"]
WOFF_ADDR_ALIASES     = ["ADDRESS STATUS"]
WOFF_CLIENT_ALIASES   = ["CLIENT STATUS"]
WOFF_UNIT_ALIASES     = ["UNIT STATUS"]
WOFF_FV_ALIASES       = ["FV REMARKS", "FIELD REMARKS"]
WOFF_REMARKS_ALIASES  = ["REMARKS", "STATUS REMARKS", "AGENCY"]
WOFF_RFD_ALIASES      = ["RFD"]


def _find_col(df: pd.DataFrame, aliases: list) -> Optional[str]:
    for alias in aliases:
        alias_up = alias.strip().upper()
        for col in df.columns:
            if alias_up in col.strip().upper():
                return col
    return None


def _pn_str(val) -> str:
    try:
        return str(int(float(str(val)))) if val and str(val).replace(".", "").isdigit() else str(val).strip()
    except Exception:
        return str(val).strip()


# ---------------------------------------------------------------------------
# Load
# ---------------------------------------------------------------------------

def load_woff_report(
    file_bytes: bytes,
    password: str = "cbs1234",
) -> tuple:
    """
    Decrypt and load the Write-Off report.
    Returns (woff_df, field_df, source_wb).
    """
    buf = decrypt_workbook(file_bytes, password)
    wb = openpyxl.load_workbook(buf, data_only=False)

    woff_sheet = next(
        (s for s in wb.sheetnames if "WOFF" in s.upper() or "WRITE" in s.upper()),
        wb.sheetnames[0],
    )
    field_sheet = next((s for s in wb.sheetnames if "FIELD" in s.upper()), None)

    buf.seek(0)
    woff_df = pd.read_excel(buf, sheet_name=woff_sheet, header=0)
    woff_df.columns = [str(c).strip() for c in woff_df.columns]
    woff_df = woff_df.dropna(how="all").reset_index(drop=True)

    buf.seek(0)
    if field_sheet:
        field_df = pd.read_excel(buf, sheet_name=field_sheet, header=0)
        field_df.columns = [str(c).strip() for c in field_df.columns]
        field_df = field_df.dropna(how="all").reset_index(drop=True)
    else:
        field_df = pd.DataFrame(columns=["PN#", "ACCOUNT NAME", "FV REMARKS"])

    return woff_df, field_df, wb


# ---------------------------------------------------------------------------
# Update
# ---------------------------------------------------------------------------

def update_woff_status(
    woff_df: pd.DataFrame,
    drr_df: Optional[pd.DataFrame],
    field_df_source: Optional[pd.DataFrame],
    cutoff_date: pd.Timestamp,
    report_date: pd.Timestamp,
) -> tuple:
    """
    Update WOFF STATUS rows from DRR and field file.
    Returns (updated_df, zero_activity_pns, changed_fv_pns, unmatched_pns).
    """
    ref = get_lookup()
    df = woff_df.copy()
    is_overall = field_df_source.attrs.get("is_overall", False) if field_df_source is not None else False

    pn_col      = _find_col(df, WOFF_PN_ALIASES) or df.columns[3]
    ac_col      = _find_col(df, WOFF_AC_ALIASES)
    acct_status = _find_col(df, WOFF_ACCT_STATUS_ALIASES)
    addr_col    = _find_col(df, WOFF_ADDR_ALIASES)
    client_col  = _find_col(df, WOFF_CLIENT_ALIASES)
    unit_col    = _find_col(df, WOFF_UNIT_ALIASES)
    fv_col      = _find_col(df, WOFF_FV_ALIASES)
    rem_col     = _find_col(df, WOFF_REMARKS_ALIASES)
    rfd_col     = _find_col(df, WOFF_RFD_ALIASES)

    zero_activity_pns: list = []
    changed_fv_pns: list    = []
    unmatched_pns: list     = []

    for idx, row in df.iterrows():
        pn = _pn_str(row.get(pn_col, ""))
        if not pn or pn.lower() == "nan":
            continue

        acct_drr = pd.DataFrame()
        if drr_df is not None and not drr_df.empty:
            acct_drr = drr_df[drr_df["account_no"] == pn]

        # STATUS REMARKS
        if drr_df is not None and rem_col:
            existing_rem = str(row.get(rem_col, "") or "")
            new_rem, had_new = update_status_remarks(existing_rem, acct_drr, cutoff_date)
            df.at[idx, rem_col] = new_rem
            if not had_new:
                zero_activity_pns.append(pn)

        # ACTION CODE + STATUS REFERENCE
        window_drr = pd.DataFrame()
        if not acct_drr.empty:
            window_drr = acct_drr[
                acct_drr["date_parsed"].notna() &
                (acct_drr["date_parsed"] >= cutoff_date)
            ]

        ac = select_daily_action_code(window_drr)
        if ac and ac_col:
            df.at[idx, ac_col] = ac

        if not window_drr.empty:
            def _rk(v):
                try:
                    return float(v) if v is not None and str(v) not in ("", "nan", "<NA>") else 999.0
                except (TypeError, ValueError):
                    return 999.0

            best_row = window_drr.copy()
            if "rank" in best_row.columns:
                best_row = best_row.sort_values("rank", key=lambda s: s.apply(_rk))
            best_status = str(best_row.iloc[0].get("action_code", "")).strip()
            ref_result = ref.lookup_drr_status(best_status)
            if ref_result.matched:
                if ref_result.client_status:
                    if acct_status:
                        df.at[idx, acct_status] = ref_result.client_status
                    if client_col:
                        df.at[idx, client_col] = ref_result.client_status
                if ref_result.address_status and addr_col:
                    df.at[idx, addr_col] = ref_result.address_status
                if ref_result.unit_status and unit_col:
                    df.at[idx, unit_col] = ref_result.unit_status
                if ref_result.rfd and rfd_col:
                    df.at[idx, rfd_col] = ref_result.rfd
            else:
                unmatched_pns.append(pn)

        # FV REMARKS
        if field_df_source is not None and not field_df_source.empty and fv_col:
            acct_field = field_df_source[field_df_source["pn"] == pn].sort_values("date_parsed", ascending=False)
            if not acct_field.empty:
                existing_fv = str(row.get(fv_col, "") or "")
                new_fv, changed = update_fv_remarks(existing_fv, acct_field, is_overall)
                df.at[idx, fv_col] = new_fv
                if changed:
                    changed_fv_pns.append(pn)

                latest_fv = acct_field.iloc[0]
                fv_text = str(latest_fv.get("FV REMARK", "") or "").strip()
                fv_ref = ref.lookup_fv_status(fv_text)
                if fv_ref.matched:
                    if fv_ref.address_status and addr_col:
                        df.at[idx, addr_col] = fv_ref.address_status
                    if fv_ref.rfd and rfd_col:
                        df.at[idx, rfd_col] = fv_ref.rfd

    return df, zero_activity_pns, changed_fv_pns, unmatched_pns


# ---------------------------------------------------------------------------
# Rebuild FIELD sheet
# ---------------------------------------------------------------------------

def rebuild_field_sheet(woff_df: pd.DataFrame) -> pd.DataFrame:
    """Rebuild the FIELD sheet: ACCOUNT NAME | PN# | ACCOUNT NAME | FV REMARKS."""
    pn_col   = _find_col(woff_df, WOFF_PN_ALIASES) or woff_df.columns[3]
    name_col = _find_col(woff_df, WOFF_NAME_ALIASES) or woff_df.columns[4]
    fv_col   = _find_col(woff_df, WOFF_FV_ALIASES)

    rows = []
    for _, row in woff_df.iterrows():
        pn   = _pn_str(row.get(pn_col, ""))
        if not pn or pn.lower() == "nan":
            continue
        name = str(row.get(name_col, "") or "").strip()
        fv   = str(row.get(fv_col, "") or "").strip() if fv_col else ""
        rows.append({"ACCOUNT NAME": name, "PN#": pn, "ACCOUNT NAME.1": name, "FV REMARKS": fv})

    return pd.DataFrame(rows)


# ---------------------------------------------------------------------------
# Write back + encrypt
# ---------------------------------------------------------------------------

def generate_woff_download(
    source_wb: openpyxl.Workbook,
    woff_df: pd.DataFrame,
    field_df: pd.DataFrame,
    password: str = "cbs1234",
) -> bytes:
    """Write updated data back into the workbook in-place and re-encrypt."""
    _write_woff_sheet(source_wb, woff_df)
    _write_field_sheet(source_wb, field_df)
    return encrypt_workbook(source_wb, password)


def _write_woff_sheet(wb: openpyxl.Workbook, df: pd.DataFrame) -> None:
    sheet_name = next(
        (s for s in wb.sheetnames if "WOFF" in s.upper() or "WRITE" in s.upper()),
        wb.sheetnames[0],
    )
    ws = wb[sheet_name]

    # Build header -> col index map (row 1 is the header)
    col_map: dict = {}
    for c in range(1, ws.max_column + 1):
        val = ws.cell(row=1, column=c).value
        if val:
            col_map[str(val).strip()] = c

    # Build PN -> excel row map
    pn_col_idx = next((cidx for hdr, cidx in col_map.items() if "PN" in hdr.upper()), None)
    if pn_col_idx is None:
        return

    pn_row_map: dict = {}
    for r in range(2, ws.max_row + 1):
        cell_val = ws.cell(row=r, column=pn_col_idx).value
        if cell_val:
            pn_row_map[_pn_str(cell_val)] = r

    pn_df_col = _find_col(df, WOFF_PN_ALIASES) or df.columns[3]

    # Columns we are allowed to update
    updatable_df_cols = [
        _find_col(df, WOFF_AC_ALIASES),
        _find_col(df, WOFF_ACCT_STATUS_ALIASES),
        _find_col(df, WOFF_ADDR_ALIASES),
        _find_col(df, WOFF_CLIENT_ALIASES),
        _find_col(df, WOFF_UNIT_ALIASES),
        _find_col(df, WOFF_FV_ALIASES),
        _find_col(df, WOFF_REMARKS_ALIASES),
        _find_col(df, WOFF_RFD_ALIASES),
    ]
    updatable_df_cols = [c for c in updatable_df_cols if c]

    for _, row in df.iterrows():
        pn = _pn_str(row.get(pn_df_col, ""))
        if not pn or pn.lower() == "nan":
            continue
        excel_row = pn_row_map.get(pn)
        if not excel_row:
            continue

        for df_col in updatable_df_cols:
            # Find matching excel column by header similarity
            excel_col = None
            df_col_up = df_col.strip().upper()[:20]
            for hdr, cidx in col_map.items():
                hdr_up = hdr.strip().upper()[:20]
                if hdr_up == df_col_up or df_col_up in hdr.upper() or hdr_up in df_col.upper():
                    excel_col = cidx
                    break
            if excel_col:
                new_val = row.get(df_col)
                if new_val is not None and str(new_val).lower() not in ("nan", "none", ""):
                    ws.cell(row=excel_row, column=excel_col).value = new_val


def _write_field_sheet(wb: openpyxl.Workbook, field_df: pd.DataFrame) -> None:
    field_sheet = next((s for s in wb.sheetnames if "FIELD" in s.upper()), None)
    if not field_sheet:
        return
    ws = wb[field_sheet]
    # Clear data rows (keep row 1 header)
    for r in range(2, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            ws.cell(row=r, column=c).value = None
    # Write new data
    cols = list(field_df.columns)
    for r_idx, (_, row) in enumerate(field_df.iterrows(), start=2):
        for c_idx, col in enumerate(cols, start=1):
            val = row.get(col)
            if val is not None and str(val).lower() not in ("nan", "none"):
                ws.cell(row=r_idx, column=c_idx).value = val
